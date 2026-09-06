#!/usr/bin/env python3
"""Build a minimal CV-like fixture docx for testing the kit scripts.

Creates scripts/tests/fixture-master.docx with the anchor structure
tailor.py expects: contact line, section header, body line, job-header
table, PERSONAL PROJECTS header, and two project blocks.
"""
import os

import docx
from docx.oxml.ns import qn

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, "fixture-master.docx")


def para_text(p):
    return "".join(t.text or "" for t in p._element.iter() if t.tag.endswith("}t"))


def build():
    doc = docx.Document()
    doc.add_paragraph("Email: test@example.com")
    doc.add_paragraph("EDUCATION")
    doc.add_paragraph("Bachelor of Test Sciences in Something")
    tbl = doc.add_table(rows=1, cols=2)
    tbl.rows[0].cells[0].text = "ACME Corp, Junior Tester"
    tbl.rows[0].cells[1].text = "2024"
    doc.add_paragraph("PERSONAL PROJECTS")
    doc.add_paragraph("Project 1 - Foo")
    doc.add_paragraph("Role: Tester")
    doc.add_paragraph("did a thing", style="List Bullet")
    doc.add_paragraph("Project 2 - Bar")
    doc.add_paragraph("Role: Tester")
    doc.add_paragraph("did another thing", style="List Bullet")

    # Explicit paragraph-level numPr on bullet paragraphs so
    # find_bullet_template() finds them (List Bullet styles carry numbering
    # at style level; tailor.py looks for it on the paragraph).
    for p in doc.paragraphs:
        if p.style.name == "List Bullet":
            ppr = p._element.get_or_add_pPr()
            numpr = p._element.makeelement(qn("w:numPr"), {})
            ilvl = p._element.makeelement(qn("w:ilvl"), {})
            ilvl.set(qn("w:val"), "0")
            numid = p._element.makeelement(qn("w:numId"), {})
            numid.set(qn("w:val"), "1")
            numpr.append(ilvl)
            numpr.append(numid)
            ppr.insert(0, numpr)

    # Ensure every paragraph carries an explicit rPr so make_para/set_text
    # always find run formatting (mirrors the font-bug lesson).
    for p in doc.paragraphs:
        runs = p.runs
        if runs:
            rpr = runs[0]._element.find(qn("w:rPr"))
            if rpr is None:
                rpr = runs[0]._element.makeelement(qn("w:rPr"), {})
                runs[0]._element.insert(0, rpr)

    doc.save(OUT)
    print(f"wrote {OUT}")


if __name__ == "__main__":
    build()
